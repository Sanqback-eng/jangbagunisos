"""매일 자동 갱신: 오늘 수집 엑셀 -> 어제와 비교 -> 사이트(docs/) 갱신.

사용법: python3 tools/daily.py <오늘 수집.xlsx> <수집일 YYYY-MM-DD>
- 오늘 수집을 data/<날짜>.json 으로 저장하고, 그 전 날짜(가장 최근) 수집과 비교해 data/drops-<날짜>.json 을 만든다.
- data/drops-*.json 전부로 날짜 페이지와 홈(최신이 위, 쌓이는 구조)을 다시 만든다.
- 파트너스 링크는 tools/links.json (상품ID:itemId 키). 없는 카드는 쿠팡 원본 주소로 두고, 링크가 필요한 목록을 links-todo-<날짜>.xlsx 로 만든다.
"""
import sys, re, json, pathlib, hashlib, datetime, shutil
import openpyxl

HERE = pathlib.Path(__file__).resolve().parent
ROOT = HERE.parent
sys.path.insert(0, str(HERE))
import build  # noqa: E402

DATA = ROOT / "data"
DOCS = ROOT / "docs"
MIN_PCT, MIN_WON = 5, 500   # 하락 인정 기준(시험값): 둘 다 만족해야 함
USE_COUPON = True           # True면 '와우회원 추가 쿠폰 N원'을 판매가에서 뺀 값을 오늘 가격으로 본다
TOP_N = 5


def num(s):
    d = re.sub(r"[^\d]", "", str(s or ""))
    return int(d) if d else None


def effective_price(sale, coupon_text, applied, wow):
    # 쿠폰 적용가 > 와우 가격 > (판매가 - 쿠폰 금액) > 판매가
    if num(applied):
        return num(applied)
    if num(wow):
        return num(wow)
    if USE_COUPON and coupon_text and "쿠폰" in str(coupon_text) and "%" not in str(coupon_text):
        c = num(coupon_text)
        if c and sale:
            return max(sale - c, 0)
    return sale


def parse(xlsx):
    """수집 엑셀 -> {키: 상품}. 같은 상품이 여러 카테고리에 있으면 먼저 나온 것만."""
    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    out = {}
    for ws in wb:
        if ws.title == "요약":
            continue
        cat = ws.title.split(" ", 1)[1]
        for r in list(ws.iter_rows(values_only=True))[1:]:
            if not r or not r[8]:
                continue
            key = ":".join(build.pkey(r[8]))
            sale = num(r[2])
            if key in out or not sale:
                continue
            out[key] = dict(
                key=key, cat=cat, rank=r[0], name=r[1], sale=sale,
                price=effective_price(sale, r[4], r[5], r[6]),
                list=num(r[3]), coupon=r[4], reviews=num(r[7]) or 0, url=r[8],
                photo=(r[9].replace("230x230ex", "492x492ex") if len(r) > 9 and r[9] else None),
            )
    return out


def prev_date(date):
    days = sorted(p.stem for p in DATA.glob("20??-??-??.json") if p.stem < date)
    return days[-1] if days else None


def compare(today, prev):
    drops = {c: [] for c in build.CATS}
    stats = dict(both=0, drop=0, rise=0, small=0, same=0)
    for k, t in today.items():
        p = prev.get(k)
        if not p:
            continue
        stats["both"] += 1
        diff = t["price"] - p["price"]
        pct = diff / p["price"] * 100
        if diff == 0:
            stats["same"] += 1
        elif diff < 0 and abs(pct) >= MIN_PCT and abs(diff) >= MIN_WON:
            stats["drop"] += 1
            drops[t["cat"]].append(dict(key=k, name=t["name"], old=p["price"], new=t["price"], pct=pct,
                                        reviews=t["reviews"], url=t["url"], photo=t["photo"]))
        elif diff > 0 and pct >= MIN_PCT and diff >= MIN_WON:
            stats["rise"] += 1
        else:
            stats["small"] += 1
    for c in drops:
        drops[c].sort(key=lambda x: x["pct"])
    return drops, stats


def title_of(date):
    d = datetime.date.fromisoformat(date)
    return f"{date}({build.WD[d.weekday()]}) 쿠팡 오늘의 가격내림 생필품"


def rebuild_site():
    css = DOCS / "style.css"
    build.STYLE_V = hashlib.md5(css.read_bytes()).hexdigest()[:8]
    links = build.load_links()
    days = []
    for p in sorted(DATA.glob("drops-*.json"), reverse=True):
        date = p.stem.replace("drops-", "")
        items = json.loads(p.read_text(encoding="utf-8"))
        for c in items:
            for it in items[c]:
                it["url"] = links.get(it["key"], it["url"])
        (DOCS / date).mkdir(parents=True, exist_ok=True)
        (DOCS / date / "index.html").write_text(build.page(date, items, title_of(date)), encoding="utf-8")
        n = sum(len(v) for v in items.values())
        mx = min([x["pct"] for v in items.values() for x in v] or [0])
        days.append((date, title_of(date), n, mx))
    (DOCS / "index.html").write_text(build.home(days), encoding="utf-8")
    return days


def links_todo(date, drops):
    links = build.load_links()
    need = [(c, it) for c in build.CATS for it in drops[c][:TOP_N] if it["key"] not in links]
    if not need:
        return None
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "링크 만들 상품"
    ws.append(["번호", "카테고리", "상품명", "쿠팡 상품 주소 (간편 링크에 붙여 넣기)", "파트너스 링크 (여기에 붙여넣기)"])
    for i, (c, it) in enumerate(need, 1):
        ws.append([i, c, it["name"], re.sub(r"&sourceType=.*", "", it["url"]), None])
    path = ROOT / f"links-todo-{date}.xlsx"
    wb.save(path)
    return path, len(need)


def main(xlsx, date):
    DATA.mkdir(exist_ok=True)
    today = parse(xlsx)
    (DATA / f"{date}.json").write_text(json.dumps(today, ensure_ascii=False), encoding="utf-8")
    pd = prev_date(date)
    print(f"오늘 {date}: 상품 {len(today)}개, 사진 {sum(1 for t in today.values() if t['photo'])}개")
    if not pd:
        print("비교할 이전 수집이 없어 오늘은 저장만 했습니다.")
        return
    prev = json.loads((DATA / f"{pd}.json").read_text(encoding="utf-8"))
    drops, st = compare(today, prev)
    (DATA / f"drops-{date}.json").write_text(json.dumps(drops, ensure_ascii=False), encoding="utf-8")
    print(f"비교 기준 {pd} / 양쪽에 있음 {st['both']} / 하락 {st['drop']} 상승 {st['rise']} 소폭 {st['small']} 동일 {st['same']}")
    print("카테고리별 하락:", {c: len(v) for c, v in drops.items()})
    rebuild_site()
    r = links_todo(date, drops)
    print("링크 필요:", f"{r[1]}개 -> {r[0].name}" if r else "없음")
    miss = [it["name"][:20] for c in build.CATS for it in drops[c][:TOP_N] if not it["photo"]]
    print("사진 없는 카드:", miss or "없음")


if __name__ == "__main__":
    main(*sys.argv[1:3])
